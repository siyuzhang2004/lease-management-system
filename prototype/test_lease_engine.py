"""Regression against the supplied Excel, including exported formula caches."""
import io
import unittest
import openpyxl
from lease_engine import baseline, calculate, audit_baseline, export_excel, COLS, SOURCE
from lease_linking import reverse_excel


class LeaseRegression(unittest.TestCase):
    def test_original_1804_values(self):
        result=calculate(baseline())
        audit=audit_baseline(result)
        self.assertEqual(audit['checked'],1804)
        self.assertEqual(audit['differences'],[])
        self.assertEqual(result['initial']['asset'],19007871.44)
        self.assertEqual(result['initial']['liability'],17422330.16)
        self.assertEqual(result['totals']['interest'],4337271.66)
        self.assertEqual(result['rows'][-1]['depreciation'],158399.96)

    def test_export_preserves_formula_values_and_structure(self):
        result=calculate(baseline());blob=export_excel(result)
        formulas=openpyxl.load_workbook(io.BytesIO(blob),data_only=False)
        cached=openpyxl.load_workbook(io.BytesIO(blob),data_only=True)
        self.assertNotIn('分录调整',formulas.sheetnames)
        s=formulas.worksheets[0];v=cached.worksheets[0]
        self.assertEqual(s['P19'].value,'=TRUNC(R18*J19/12,2)')
        self.assertEqual(s['U138'].value,'=T138')
        self.assertEqual(s['C6'].value,'=计算参数!B7')
        self.assertEqual(v['C6'].value.date().isoformat(),'2033-08-03')
        for row in result['rows']:
            for col,key in COLS.items():self.assertAlmostEqual(v[f'{col}{row["period"]+18}'].value,row[key],places=7)
        self.assertEqual(cached['租金参数']['F5'].value,1585541.28)
        self.assertEqual(cached['年度汇总']['C2'].value,307032.46)
        self.assertTrue(formulas['年度汇总']['C2'].value.startswith('=SUM('))
        self.assertIsNone(s['G7'].value)
        self.assertIsNone(s['V8'].value)
        self.assertEqual(v['F141'].value,v['D141'].value)
        self.assertEqual(v['T18'].value,19007871.44)
        self.assertEqual(s.freeze_panes,'D19')

    def test_reverse_original_with_provenance(self):
        report=reverse_excel(SOURCE.read_bytes())
        self.assertTrue(report['comparison']['matched'])
        self.assertEqual(report['comparison']['checked'],1804)
        self.assertEqual(report['params']['availableDate'],'')
        self.assertEqual(len(report['warnings']),2)
        self.assertEqual(report['params']['stages'],baseline()['stages'])

    def test_linked_export_restores_identity_and_changed_parameters(self):
        p=baseline();p['recordId']='LEASE-TEST';p['revision']=7
        p['stages'][1]['gross']+=100000;p['displayMode']='auto'
        blob=export_excel(calculate(p));report=reverse_excel(blob)
        self.assertTrue(report['comparison']['matched'])
        self.assertEqual(report['comparison']['checked'],1804)
        self.assertEqual(report['params']['recordId'],'LEASE-TEST')
        self.assertEqual(report['params']['revision'],7)
        self.assertEqual(report['params']['stages'][1]['gross'],p['stages'][1]['gross'])
        book=openpyxl.load_workbook(io.BytesIO(blob),data_only=False)
        self.assertEqual(book['租金参数']['E6'].value,'=计算参数!$B$10')
        self.assertEqual(book['租赁计算明细']['F30'].hyperlink.location,"'租金参数'!D6")
        self.assertEqual(book['租金参数']['D6'].hyperlink.location,"'租赁计算明细'!F30")

    def test_edited_excel_without_recalculation_not_claimed_matched(self):
        book=openpyxl.load_workbook(io.BytesIO(export_excel(calculate(baseline()))))
        book['租金参数']['D6']=2000000
        stream=io.BytesIO();book.save(stream)
        # Saving through openpyxl removes formula caches. Do not mistake these for zeros.
        report=reverse_excel(stream.getvalue())
        self.assertFalse(report['comparison']['matched'])
        self.assertGreater(report['comparison']['missing'],0)
        self.assertEqual(report['params']['stages'][1]['gross'],2000000)

    def test_changed_amount_updates_present_value_and_closes(self):
        p=baseline();p['stages'][1]['gross']+=100000;p['displayMode']='auto'
        result=calculate(p)
        self.assertGreater(result['initial']['liability'],17422330.16)
        self.assertEqual(result['rows'][-1]['closingLiability'],0)
        self.assertEqual(result['rows'][-1]['closingAsset'],0)
        self.assertAlmostEqual(sum(r['monthlyRent'] for r in result['rows']),result['initial']['totalFixed'],places=5)

    def test_zero_rate_arrears(self):
        p=baseline();p.update(months=12,end='2024-08-03',annualRate=0,displayMode='auto')
        p['stages']=[dict(months=12,gross=120000,paymentPeriod=12)]
        result=calculate(p)
        self.assertEqual(result['initial']['upfront'],0)
        self.assertEqual(result['totals']['interest'],0)
        self.assertEqual(result['rows'][-1]['closingLiability'],0)

    def test_advanced_terms_preserved_with_contract_export(self):
        p=baseline();p['advanced']={'rateEvidence':'合同利率判断依据','variableType':'other','otherEstimate':'5000','otherFormula':'按销售额结算'}
        result=calculate(p);report=reverse_excel(export_excel(result))
        self.assertEqual(report['params']['advanced'],p['advanced'])
        self.assertTrue(report['comparison']['matched'])
        self.assertEqual(result['initial']['liability'],17422330.16)

    def test_unsupported_advanced_cashflow_cannot_be_silently_ignored(self):
        for terms in [{'variableType':'index','indexMonthly':'1000'},{'renewalEnabled':True,'renewalDecision':'yes'},
                      {'substanceEnabled':True,'substanceMonthly':'1000'},{'residualEnabled':True,'residualAmount':'5000'}]:
            p=baseline();p['advanced']=terms
            with self.assertRaisesRegex(ValueError,'高级参数已保存'):calculate(p)

    def test_invalid_periods_rejected(self):
        p=baseline();p['stages'][0]['paymentPeriod']=121
        with self.assertRaises(ValueError):calculate(p)

    def test_arbitrary_phase_lengths_and_payment_frequencies(self):
        p=baseline();p.update(months=12,end='2024-08-03',taxRate=0,displayMode='auto')
        p['rentPhases']=[dict(months=5,gross=5000,frequency=3,timing='advance'),dict(months=7,gross=14000,frequency=3,timing='arrears')]
        r=calculate(p)
        self.assertEqual([s['months'] for s in r['stages']],[3,2,3,3,1])
        self.assertEqual([(s['paymentPeriod'],s['gross']) for s in r['stages']],[(0,3000),(3,2000),(8,6000),(11,6000),(12,2000)])
        self.assertEqual(r['rows'][-1]['closingLiability'],0)
        report=reverse_excel(export_excel(r))
        self.assertTrue(report['comparison']['matched'])
        self.assertEqual([s['months'] for s in report['params']['rentPhases']],[5,7])

    def test_month_beginning_versus_month_end_changes_initial_liability(self):
        p=baseline();p.update(months=18,end='2025-02-03',taxRate=0,displayMode='auto')
        p['rentPhases']=[dict(months=18,gross=18000,frequency=1,timing='advance')]
        advance=calculate(p);p['rentPhases'][0]['timing']='arrears';arrears=calculate(p)
        self.assertEqual(advance['initial']['upfront'],1000)
        self.assertEqual(arrears['initial']['upfront'],0)
        self.assertGreater(arrears['initial']['liability'],advance['initial']['liability'])
        self.assertEqual(advance['rows'][-1]['closingLiability'],0)
        self.assertEqual(arrears['rows'][-1]['closingLiability'],0)

    def test_phase_migration_preserves_original_1804_values(self):
        p=baseline();p['rentPhases']=[dict(s,frequency=0,timing='advance') for s in p['stages']]
        r=calculate(p);self.assertTrue(audit_baseline(r)['matched'])
        self.assertTrue(reverse_excel(export_excel(r))['comparison']['matched'])

    def test_phase_length_does_not_extend_contract(self):
        p=baseline();p['rentPhases']=[dict(months=121,gross=1000,frequency=1,timing='advance')]
        with self.assertRaisesRegex(ValueError,'不会改变合同期限'):calculate(p)
        self.assertEqual(p['end'],'2033-08-03')
        p=baseline();p['months']=119
        with self.assertRaises(ValueError):calculate(p)


if __name__=='__main__':unittest.main()
