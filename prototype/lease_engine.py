"""Local lease calculation and Excel export, following the supplied DTT workbook.

Annual rate / 12, per-payment PV truncation, monthly interest truncation,
last-payment interest adjustment and last-period depreciation true-up are deliberate.
"""
from __future__ import annotations

import calendar
import copy
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_DOWN, ROUND_HALF_UP
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / '西安公司使用权资产.xlsx'


def cents(value):
    return float(Decimal(str(value)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))


def trunc(value):
    return float(Decimal(str(value)).quantize(Decimal('.01'), rounding=ROUND_DOWN))


def shift_month(iso, n):
    d = date.fromisoformat(iso)
    i = d.year * 12 + d.month - 1 + n
    y, m = divmod(i, 12)
    return date(y, m + 1, min(d.day, calendar.monthrange(y, m + 1)[1]))


def baseline():
    gross = [1728240] * 3 + [2720035] * 2 + [2831057] * 2 + [2997590] * 2 + [3164122]
    monthly = [132128.44] * 3 + [207953.75] * 2 + [216441.66] * 2 + [229173.52] * 2 + [241905.39]
    tails = [.05] * 3 + [-.13] * 2 + [.05] * 2 + [.25] * 2 + [-.23]
    return {
        'recordId': 'LS-XA-2023-001', 'contractId': 'LS-XA-2023-001', 'revision': 1,
        'company': '西安普源',
        'name': '秦创原集成电路加速器（西安电子谷核心区）国家电子实验园H区5#',
        'lessor': '西安高新丝路通信创新谷有限公司',
        'availableDate': '2023-05-04', 'freeEnd': '2023-08-03',
        'start': '2023-08-04', 'end': '2033-08-03', 'months': 120,
        'annualRate': .042, 'taxRate': .09, 'area': 4625.91, 'deposit': 432060,
        'assetAdjustment': 0, 'displayMode': 'source',
        'policy': '公司简便口径：初始装修免租期不纳入确认及折旧；按德勤原表从2023-08-04起计算。',
        'stages': [dict(months=12, gross=g, paymentPeriod=i*12,
                        monthlyNet=monthly[i], monthlyTail=tails[i]) for i, g in enumerate(gross)],
    }


def validate(p):
    errors = []
    advanced=p.get('advanced') or {}
    if not isinstance(advanced,dict):raise ValueError('高级参数格式不正确。')
    for key in ('substanceMonthly','indexMonthly','otherEstimate','purchasePrice','terminationPenalty','residualAmount','residualCap','incentiveAmount','prepaid','directCosts','restorationPV','serviceMonthly'):
        value=float(advanced.get(key) or 0)
        if not math.isfinite(value) or value<0:errors.append('高级参数金额须为有效非负数：'+key)
    pending=[]
    if advanced.get('freeEnabled'):
        free_periods=advanced.get('freePeriods') or []
        if len(free_periods)>1:pending.append('后续免租区间请折算到对应租金阶段')
        if any(not r.get('start') or not r.get('end') or r['end']<r['start'] for r in free_periods):errors.append('请完整填写有效的免租期起止日期。')
    if advanced.get('stepMode','none')!='none':pending.append('固定递增条款请换算到分阶段租金，并将递增方式设为不重复递增')
    if advanced.get('substanceEnabled') and float(advanced.get('substanceMonthly') or 0)>0:pending.append('额外实质固定付款需补充对应现金流')
    if advanced.get('variableType') in ('index','both'):pending.append('指数或比率挂钩付款需按开始日水平补充现金流')
    for key,label in [('renewal','续租'),('purchase','购买'),('termination','提前终止')]:
        if advanced.get(key+'Enabled') and advanced.get(key+'Decision') not in ('no',):pending.append(label+'选择权需完成判断并补充相应期间或付款')
    if advanced.get('residualEnabled') and float(advanced.get('residualAmount') or 0)>0:pending.append('担保余值需补充预计支付时点')
    if advanced.get('incentiveEnabled') and not advanced.get('incentiveReceived') and float(advanced.get('incentiveAmount') or 0)>0:pending.append('未来租赁激励需补充抵减现金流')
    if float(advanced.get('prepaid') or 0)>0:pending.append('预付租金需匹配已有约定付款，避免重复计入')
    if advanced.get('servicePolicy')=='combine' and float(advanced.get('serviceMonthly') or 0)>0:pending.append('合并服务费需补充对应合同现金流')
    if advanced.get('taxBasis')=='review':pending.append('请先复核可抵扣增值税口径')
    if advanced.get('exemption','none')!='none':pending.append('简化处理需另行复核，不能沿用一般租赁计算表')
    if pending:errors.append('高级参数已保存，尚未生成新版计算表：'+'；'.join(pending)+'。')
    for key in ('company', 'name', 'lessor', 'start', 'end'):
        if not str(p.get(key, '')).strip(): errors.append(f'请填写{key}。')
    try:
        if p.get('freeStart') and p.get('freeEnd') and date.fromisoformat(p['freeStart'])>date.fromisoformat(p['freeEnd']):errors.append('免租期开始日不得晚于结束日。')
        start, end = date.fromisoformat(p['start']), date.fromisoformat(p['end'])
        if end < start: errors.append('租赁结束日不得早于租赁开始日。')
    except (ValueError, KeyError): errors.append('请填写有效的起止日期。')
    n = p.get('months', 0)
    if not isinstance(n, int) or not 1 <= n <= 600: errors.append('计算期数须为1至600的整数月。')
    elif not errors:
        expected_end=shift_month(p['start'], n)-timedelta(days=1)
        if end != expected_end: errors.append(f'租赁结束日与阶段月数不一致；按当前起算日及{n}期，应为{expected_end.isoformat()}。')
    try:
        if p.get('freeEnd') and not p.get('recognizeDuringFreePeriod') and date.fromisoformat(p['freeEnd']) >= date.fromisoformat(p['start']):
            errors.append('公司简便口径下，租赁开始日须晚于免租期结束日。')
        if p.get('recognizeDuringFreePeriod') and p.get('freeEnd'):
            if not p.get('freeStart') or not p['start'] <= p['freeStart'] <= p['freeEnd'] <= p['end']:
                errors.append('纳入确认期间的免租期须完整填写并位于租赁期内。')
        if p.get('availableDate') and date.fromisoformat(p['availableDate']) > date.fromisoformat(p['start']):
            errors.append('取得资产控制权日不得晚于租赁开始日。')
    except (ValueError, KeyError): errors.append('取得资产控制权日或免租期结束日格式不正确。')
    for key, maximum in [('annualRate', 1), ('taxRate', 1), ('area', 1e9), ('deposit', 1e12)]:
        val = p.get(key, 0)
        if not isinstance(val, (float, int)) or not math.isfinite(val) or not 0 <= val <= maximum:
            errors.append(f'{key}超出允许范围。')
    if not isinstance(p.get('assetAdjustment', 0), (float, int)) or not math.isfinite(p.get('assetAdjustment', 0)):
        errors.append('使用权资产成本调整须为有效金额。')
    stages = p.get('stages', [])
    if not isinstance(stages, list) or not 1 <= len(stages) <= 600:
        errors.append('请填写租金阶段。'); stages = []
    total = 0
    for i, s in enumerate(stages):
        count = s.get('months', 0)
        if not isinstance(count, int) or count < 1: errors.append(f'第{i+1}段月数须为正整数。')
        else: total += count
        for key in ('gross', 'monthlyNet'):
            value = s.get(key, 0)
            if not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0:
                errors.append(f'第{i+1}段金额须为有效的非负数。')
        value = s.get('monthlyTail', 0)
        if not isinstance(value, (float, int)) or not math.isfinite(value): errors.append(f'第{i+1}段尾差须为有效金额。')
        period = s.get('paymentPeriod')
        if not isinstance(period, int) or not 0 <= period <= n: errors.append(f'第{i+1}段支付期须为0至{n}的整数。')
    if total != n: errors.append(f'阶段月数合计{total}，须等于计算期数{n}。')
    if errors: raise ValueError('\n'.join(errors))


def expand_rent_phases(p):
    phases=p.get('rentPhases')
    if phases is None:return
    if not isinstance(phases,list) or not 1<=len(phases)<=600:raise ValueError('请填写至少一个租金阶段。')
    stages=[];offset=0
    for index,phase in enumerate(phases):
        count=phase.get('months');gross=phase.get('gross');frequency=phase.get('frequency',0);timing=phase.get('timing')
        if not isinstance(count,int) or isinstance(count,bool) or not 1<=count<=600:raise ValueError(f'第{index+1}阶段请填写1至600之间的整数月数。')
        if not isinstance(gross,(int,float)) or not math.isfinite(gross) or gross<0:raise ValueError(f'第{index+1}阶段请填写有效的租金总额。')
        if not isinstance(frequency,int) or not 0<=frequency<=600:raise ValueError(f'第{index+1}阶段付款频率不正确。')
        if timing not in ('advance','arrears','custom'):raise ValueError(f'第{index+1}阶段请选择期初付或期末付。')
        expected_end=(shift_month(p['start'],offset+count)-timedelta(days=1)).isoformat()
        if phase.get('endDate') and phase['endDate']!=expected_end:raise ValueError(f'第{index+1}阶段结束日须与整月月数一致；当前应为{expected_end}。')
        frequency=frequency or count
        if timing=='custom' and frequency<count:raise ValueError('原约定期次只适用于该阶段一次支付；分次付款请选择期初付或期末付。')
        allocated=Decimal('0')
        for pos in range(0,count,frequency):
            chunk=min(frequency,count-pos)
            amount=cents(Decimal(str(gross))*chunk/count) if pos+chunk<count else cents(Decimal(str(gross))-allocated)
            allocated+=Decimal(str(amount))
            period=phase.get('paymentPeriod') if timing=='custom' else offset+pos+(chunk if timing=='arrears' else 0)
            stages.append(dict(months=chunk,gross=amount,paymentPeriod=period,sourcePhase=index,
                               monthlyNet=phase.get('monthlyNet',0),monthlyTail=phase.get('monthlyTail',0) if pos+chunk==count else 0))
        offset+=count
    if offset!=p['months']:raise ValueError(f'租金阶段合计{offset}个月，合同租赁期为{p["months"]}个月。请调整阶段月数；增删阶段不会改变合同期限。')
    p['stages']=stages


def calculate(params):
    p = copy.deepcopy(params)
    precision = p.get('calculationPrecision', 2)
    if precision not in (0, 2):
        raise ValueError('计算精度只能为整数或两位小数。')
    def truncate_amount(value):
        return float(Decimal(str(value)).quantize(Decimal('1') if precision == 0 else Decimal('.01'), rounding=ROUND_DOWN))
    rounding = p.get('roundingPolicy', 'truncate-payments')
    if rounding not in ('truncate-payments', 'round-total-pv'):
        raise ValueError('请选择有效的金额舍入口径。')
    amount_round = cents if rounding == 'round-total-pv' else truncate_amount
    p.setdefault('recordId','LEASE-'+hashlib.sha256((str(p.get('company'))+'|'+str(p.get('name'))+'|'+str(p.get('start'))).encode()).hexdigest()[:12])
    p.setdefault('revision',1)
    if not isinstance(p['recordId'],str) or not 1<=len(p['recordId'])<=80 or not isinstance(p['revision'],int) or isinstance(p['revision'],bool) or p['revision']<1:
        raise ValueError('关联记录编号或版本不正确。')
    if p.get('calendarTerms') is not None:
        from calendar_contract import calculate_calendar_contract
        return calculate_calendar_contract(p)
    if p.get("paymentSchedule") is not None:
        from scheduled_lease import calculate_scheduled
        return calculate_scheduled(p)
    expand_rent_phases(p)
    validate(p)
    n, annual, tax = p['months'], p['annualRate'], p['taxRate']
    rate = annual / 12
    payments = [0.] * (n + 1)
    monthly = []
    stages = []
    for i, stage in enumerate(p['stages']):
        count = stage['months']
        net = cents(Decimal(str(stage['gross'])) / (1 + Decimal(str(tax))))
        base = stage.get('monthlyNet', 0) if p.get('displayMode') == 'source' else cents(Decimal(str(net)) / count)
        tail = stage.get('monthlyTail', 0) if p.get('displayMode') == 'source' else cents(Decimal(str(net)) - Decimal(str(base)) * count)
        stages.append({**stage, 'net': net, 'monthlyNet': base, 'monthlyTail': tail,
                       'fromPeriod': len(monthly)+1, 'toPeriod': len(monthly)+count})
        monthly.extend([base] * (count - 1) + [cents(Decimal(str(base)) + Decimal(str(tail)))])
        payments[stage['paymentPeriod']] = cents(Decimal(str(payments[stage['paymentPeriod']])) + Decimal(str(net)))
    if p.get('recognizeDuringFreePeriod') and p.get('freeEnd'):
        for i, amount in enumerate(monthly):
            begin = shift_month(p['start'], i).isoformat()
            if p['freeStart'] <= begin <= p['freeEnd'] and amount:
                raise ValueError('免租月份的固定租金须为零，请同步调整租金阶段。')
    pv = [0.] + [amount_round(payments[i] / (1+rate)**i) for i in range(1, n+1)]
    initial_liability = cents(sum(Decimal(str(v)) for v in pv))
    if rounding == 'round-total-pv':
        initial_liability = cents(sum(Decimal(str(payments[i])) / (1+Decimal(str(annual))/12)**i for i in range(1,n+1)))
    initial_asset = cents(Decimal(str(initial_liability)) + Decimal(str(payments[0])) + Decimal(str(p.get('assetAdjustment', 0))))
    if initial_asset < 0: raise ValueError('调整后的使用权资产成本不能为负。')
    total_fixed = cents(sum(Decimal(str(s['net'])) for s in stages))
    final_payment = max((i for i, value in enumerate(payments) if i and value), default=0)
    liability, asset, cumulative = initial_liability, initial_asset, payments[0]
    depreciation = amount_round(Decimal(str(initial_asset)) / n)
    rows = []
    start = date.fromisoformat(p['start'])
    for i in range(1, n+1):
        begin = start if i == 1 else date(shift_month(p['start'], i-1).year, shift_month(p['start'], i-1).month, 1)
        # The original table uses 120 rows, with its final row ending on 2033-08-03.
        end = date.fromisoformat(p['end']) if i == n else date(begin.year, begin.month, calendar.monthrange(begin.year, begin.month)[1])
        if begin > end: raise ValueError('起止日期与计算期数不匹配，末期开始日晚于结束日。')
        interest = cents(Decimal(str(payments[i])) - Decimal(str(liability))) if i == final_payment else amount_round(Decimal(str(liability))*Decimal(str(annual))/12)
        principal = cents(Decimal(str(payments[i])) - Decimal(str(interest)))
        closing_liability = cents(Decimal(str(liability)) - Decimal(str(principal)))
        dep = asset if i == n else depreciation
        closing_asset = cents(Decimal(str(asset)) - Decimal(str(dep)))
        cumulative = cents(Decimal(str(cumulative)) + Decimal(str(payments[i])))
        rows.append(dict(period=i, start=begin.isoformat(), end=end.isoformat(), monthlyRent=monthly[i-1],
                         payment=payments[i], cumulativePayment=cumulative, remainingPayments=cents(Decimal(str(total_fixed))-Decimal(str(cumulative))),
                         annualRate=annual, pvFactor=1/(1+rate)**i, openingLiability=liability,
                         discountedPayment=pv[i], principal=principal, interest=interest, totalPayment=payments[i],
                         closingLiability=closing_liability, openingAsset=asset, depreciation=dep, closingAsset=closing_asset))
        liability, asset = closing_liability, closing_asset
    years = {}
    for r in rows:
        y = r['end'][:4]
        item = years.setdefault(y, dict(year=y, payment=0, interest=0, depreciation=0))
        for key in ('payment', 'interest', 'depreciation'): item[key] = cents(Decimal(str(item[key])) + Decimal(str(r[key])))
        item.update(asset=r['closingAsset'], liability=r['closingLiability'])
    return dict(params=p, stages=stages, rows=rows, annual=list(years.values()),
                initial=dict(asset=initial_asset, liability=initial_liability, upfront=payments[0], totalFixed=total_fixed),
                totals=dict(payment=total_fixed, interest=cents(sum(Decimal(str(r['interest'])) for r in rows)),
                            depreciation=cents(sum(Decimal(str(r['depreciation'])) for r in rows))),
                notes=[
                    '租赁负债按尚未支付租赁付款额的现值初始计量。',
                    '租赁付款按照各阶段约定的金额、付款频率及期初或期末支付安排生成。',
                    '利息按年折现率折算至各期计提；末次付款及折旧对尾差作结清处理。',
                    f'本表共{n}期，租赁开始日为{p["start"]}，租赁结束日为{p["end"]}。',
                    '本表反映合同约定的使用权资产与租赁负债余额，不记录实际付款状态。',
                    '租赁负债期末余额＝期初余额＋当期利息－当期约定付款额。',

                ])


COLS = {'D':'monthlyRent','F':'payment','G':'cumulativePayment','H':'remainingPayments','J':'annualRate',
        'K':'pvFactor','M':'openingLiability','N':'discountedPayment','O':'principal','P':'interest',
        'Q':'totalPayment','R':'closingLiability','T':'openingAsset','U':'depreciation','V':'closingAsset'}


def audit_baseline(result=None):
    result = result or calculate(baseline())
    s = openpyxl.load_workbook(SOURCE, data_only=True).worksheets[0]
    failures, checked = [], 0
    for row in result['rows']:
        for col, key in COLS.items():
            address = f'{col}{row["period"]+18}'
            expected = s[address].value or 0
            diff = abs(float(expected)-row[key])
            tolerance = 1e-12 if key in ('annualRate', 'pvFactor') else .00001
            checked += 1
            if diff > tolerance: failures.append(dict(cell=address, expected=expected, actual=row[key], difference=diff))
    for address, key in [('R18','liability'),('T18','asset'),('F18','upfront'),('C7','totalFixed')]:
        expected = s[address].value
        checked += 1
        if abs(expected-result['initial'][key]) > .00001: failures.append(dict(cell=address,expected=expected,actual=result['initial'][key]))
    return dict(checked=checked, differences=failures, matched=not failures)


def _cache_xlsx(blob, caches):
    ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as src, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
        for entry in src.infolist():
            data = src.read(entry.filename)
            if entry.filename in caches:
                root = ET.fromstring(data)
                for c in root.findall('.//s:c', ns):
                    value = caches[entry.filename].get(c.attrib['r'])
                    if value is None or c.find('s:f',ns) is None: continue
                    old = c.find('s:v',ns)
                    if old is None: old=ET.SubElement(c, '{'+ns['s']+'}v')
                    old.text=str(value)
                    if isinstance(value,str):c.set('t','str')
                data=ET.tostring(root, encoding='utf-8', xml_declaration=True)
            dst.writestr(entry, data)
    return out.getvalue()


def export_excel(result):
    p, rows, stages = result['params'], result['rows'], result['stages']
    w = openpyxl.load_workbook(SOURCE, keep_links=False)
    s = w.worksheets[0]
    for other in w.worksheets[1:]: w.remove(other)
    s.title = '租赁计算明细'
    w.defined_names.clear()
    # Retain the supplied main table and styling, omit its SAP voucher/audit-link blocks.
    for merged in list(s.merged_cells.ranges):
        if merged.max_col > 22: s.unmerge_cells(str(merged))
    for row in s:
        for c in row:
            if c.column > 22 or c.row >= 18: c.value = None
    s._images=[]
    s.delete_cols(23, max(0,s.max_column-22))
    for address in ('B8','B9','G2','G7','V1','V8','V9','V10'):s[address]=None
    s.merge_cells('A8:B8');s.merge_cells('A9:B9')
    n = p['months']; last=18+n; total=last+3; end_stage=4+len(stages)
    side = Side(style='thin',color='D8DEE5')
    blue = PatternFill('solid',fgColor='E8EEF5'); yellow=PatternFill('solid',fgColor='FFF8CE')
    cache = {}
    def put(address, value, cached=None):
        s[address]=value
        if cached is not None: cache[address]=cached
    for address,value in {'C1':p['company'],'C2':p['name'],'C3':p['lessor'],'C4':date.fromisoformat(p['start']),
                          'C5':f'{n}个月','C6':date.fromisoformat(p['end']),'C10':p.get('area',0),'C11':p.get('deposit',0)}.items(): put(address,value)
    for main_cell,param_cell,value in [('C1','B2',p['company']),('C2','B3',p['name']),('C3','B4',p['lessor']),
                                       ('C4','B6',to_excel(date.fromisoformat(p['start']))),('C6','B7',to_excel(date.fromisoformat(p['end']))),
                                       ('C10','B11',p.get('area',0)),('C11','B12',p.get('deposit',0))]:
        put(main_cell,f'=计算参数!{param_cell}',value)
    put('C5','=计算参数!B8&"个月"',f'{n}个月')
    put('C7',f'=F{total}',result['initial']['totalFixed']);put('C8','=计算参数!B9',p['annualRate']);put('C9','=C8/12',p['annualRate']/12)
    s['A4']='租赁开始日';s['A5']='计算期数';s['A6']='租赁结束日';s['A8']='年利率（实际利率法）';s['A9']='月计息系数（年÷12）';s['F15']='合同约定支付';s['F16']='本期';s['M16']='支付前余额';s['O16']='净减少额（付款－利息）';s['R16']='支付后余额';s['T16']='折旧前余额'
    s['H14']='合同付款额';s['H15']='未折现金额';s['H16']='剩余合同现金流';s['M14']='租赁负债余额';s['R14']='期末租赁负债';s['T14']='使用权资产摊销'
    for key in ('C1','C2','C3'):
        s[key].alignment=Alignment(vertical='center',wrap_text=True)
        s[key].font=Font(name='微软雅黑',size=11,color='27313D')
    s.row_dimensions[2].height=60;s.row_dimensions[3].height=32
    for row in (8,9):
        s.row_dimensions[row].height=32
        s[f'A{row}'].alignment=Alignment(vertical='center',wrap_text=True)
        s[f'A{row}'].font=Font(name='微软雅黑',size=11,color='27313D')
    s['N16']='现值计算';s['J16']='年利率';s['S16']=''
    for row in range(18,total+1):
        style_row=18 if row==18 else 19
        for col in range(1,23):
            c=s.cell(row,col); ref=s.cell(style_row,col)
            if row != style_row: c._style=copy.copy(ref._style)
            c.font=Font(name='微软雅黑',size=10,color='27313D')
            c.alignment=Alignment(horizontal='right' if col>=4 else 'center',vertical='center')
            c.number_format='0.00%' if col==10 else '0.000000000' if col==11 else '#,##0.00;[Red](#,##0.00);–'
            if col in (2,3):c.number_format='yyyy/mm/dd'
        s.row_dimensions[row].height=21
    sf=f"'租金参数'!$F$5:$F${end_stage}";sp=f"'租金参数'!$I$5:$I${end_stage}"
    put('F18',f'=SUMIF({sp},0,{sf})',result['initial']['upfront']);put('G18','=F18',result['initial']['upfront'])
    put('H18',f'=$C$7-G18',cents(result['initial']['totalFixed']-result['initial']['upfront']))
    put('R18',f'=N{total}',result['initial']['liability']);put('T18','=R18+F18+计算参数!B13',result['initial']['asset']);put('V18','=T18',result['initial']['asset'])
    s['B18']='初始确认';s.merge_cells('B18:D18')
    final_pay=max((r['period'] for r in rows if r['payment']),default=0)
    for r in rows:
        i=r['period'];row=i+18;prev=row-1
        stage_id=next(k for k,st in enumerate(stages) if st['fromPeriod']<=i<=st['toPeriod']); sr=stage_id+5
        s[f'A{row}']=i;s[f'A{row}'].number_format='0'
        put(f'B{row}','=计算参数!B6' if i==1 else f'=C{prev}+1',to_excel(date.fromisoformat(r['start'])))
        put(f'C{row}','=计算参数!B7' if i==n else f'=EOMONTH(B{row},0)',to_excel(date.fromisoformat(r['end'])))
        formulas={
            'D':f"='租金参数'!G{sr}"+(f"+'租金参数'!H{sr}" if i==stages[stage_id]['toPeriod'] else ''),
            'F':f'=SUMIF({sp},A{row},{sf})','G':f'=G{prev}+F{row}','H':f'=$C$7-G{row}',
            'J':'=$C$8','K':f'=1/(1+J{row}/12)^A{row}',
            'M':f'=R{prev}','N':f'=TRUNC(F{row}*K{row},2)','O':f'=F{row}-P{row}',
            'P':f'=F{row}-R{prev}' if i==final_pay else f'=TRUNC(R{prev}*J{row}/12,2)',
            'Q':f'=SUM(O{row}:P{row})','R':f'=R{prev}-O{row}','T':f'=V{prev}',
            'U':f'=T{row}' if i==n else f'=TRUNC($V$18/$A${last},2)','V':f'=T{row}-U{row}',
        }
        for col,formula in formulas.items():put(f'{col}{row}',formula,r[COLS[col]])
    s[f'B{total}']='合计'
    for col,start,key in [('D',19,None),('F',18,None),('N',19,None),('O',19,None),('P',19,None),('Q',19,None),('U',19,None)]:
        value = result['initial']['totalFixed'] if col=='F' else cents(sum(Decimal(str(r[COLS[col]])) for r in rows))
        put(f'{col}{total}',f'=ROUND(SUM({col}{start}:{col}{last}),2)',value)
    put(f'K{total}',f'=K{last}',rows[-1]['pvFactor'])
    for c in s[total][:22]: c.fill=blue;c.font=Font(name='微软雅黑',size=10,bold=True,color='243040')
    s.freeze_panes='D19';s.sheet_view.zoomScale=75;s.sheet_view.showGridLines=False
    s.print_area=f'A1:V{total}';s.print_title_rows='12:16';s.sheet_properties.pageSetUpPr.fitToPage=True
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    s.sheet_properties.outlinePr.summaryRight=False
    for col in ('E','I','L','S'):s.column_dimensions[col].hidden=True
    s.column_dimensions['A'].width=7;s.column_dimensions['B'].width=13;s.column_dimensions['C'].width=13
    for col in ('D','F','G','H','M','N','O','P','Q','R','T','U','V'):s.column_dimensions[col].width=18
    for row in range(12,17):s.row_dimensions[row].height=24

    ps=w.create_sheet('计算参数')
    items=[('计算依据','西安公司德勤逻辑表 / 本地生成'),('公司',p['company']),('租赁项目',p['name']),('出租方',p['lessor']),
           ('取得资产控制权日',p.get('availableDate','')),('租赁开始日',p['start']),('租赁结束日',p['end']),('计算期数',n),
           ('年利率',p['annualRate']),('增值税率',p['taxRate']),('面积（平方米）',p.get('area',0)),('押金（单独记录）',p.get('deposit',0)),
           ('初始直接费用等（合计）',p.get('assetAdjustment',0)),('免租期结束日',p.get('freeEnd','')),('计算口径',p.get('policy','')),
           ('月租金展示模式','保留原表数值及尾差' if p.get('displayMode')=='source' else '按阶段净额自动均摊'),
           ('所属合同编号',p['recordId']),('参数与计算表版本',p['revision']),('关联说明','本合同参数驱动本合同计算表；点击明细可返回参数。期限、阶段数量变动请回系统重新生成。'),('免租期开始日',p.get('freeStart',p.get('availableDate','') if p.get('freeEnd') else ''))]
    for item in items:ps.append(item)
    for cell in ('B5','B6','B7','B14','B20'):
        if ps[cell].value:ps[cell]=date.fromisoformat(str(ps[cell].value));ps[cell].number_format='yyyy/mm/dd'
    ps['B9'].number_format=ps['B10'].number_format='0.00%'
    rs=w.create_sheet('租金参数')
    rs.append(['固定付款分阶段参数']);rs.append(['第0期表示确认时点支付；其他期数对应主表期数。金额均为合同约定现金流。'])
    rs.append([]);rs.append(['阶段','自第几期','至第几期','含税阶段租金','税率','不含税阶段租金','月租金展示值','段末展示尾差','约定支付期','原表月租金输入','原表尾差输入'])
    rent_cache={}
    for i,st in enumerate(stages,5):
        rs.append([i-4,st['fromPeriod'],st['toPeriod'],st['gross'],'=计算参数!$B$10',f'=ROUND(D{i}/(1+E{i}),2)',
                   f'=IF(LEFT(计算参数!$B$16,2)="保留",J{i},ROUND(F{i}/(C{i}-B{i}+1),2))',
                   f'=IF(LEFT(计算参数!$B$16,2)="保留",K{i},ROUND(F{i}-G{i}*(C{i}-B{i}+1),2))',
                   st['paymentPeriod'],p['stages'][i-5].get('monthlyNet',0),p['stages'][i-5].get('monthlyTail',0)])
        rent_cache[f'E{i}']=p['taxRate'];rent_cache[f'F{i}']=st['net'];rent_cache[f'G{i}']=st['monthlyNet'];rent_cache[f'H{i}']=st['monthlyTail'];rs[f'E{i}'].number_format='0.00%'
    ys=w.create_sheet('年度汇总');ys.append(['年度','合同租金现金流（不含初始支付）','利息费用','使用权资产折旧','期末租赁负债','期末使用权资产'])
    annual_cache={}
    for index,a in enumerate(result['annual'],2):
        periods=[r['period']+18 for r in rows if r['end'].startswith(a['year'])];first,last_in_year=min(periods),max(periods)
        ys.append([int(a['year']),f"=SUM('租赁计算明细'!F{first}:F{last_in_year})",f"=SUM('租赁计算明细'!P{first}:P{last_in_year})",
                   f"=SUM('租赁计算明细'!U{first}:U{last_in_year})",f"='租赁计算明细'!R{last_in_year}",f"='租赁计算明细'!V{last_in_year}"])
        for col,key in [('B','payment'),('C','interest'),('D','depreciation'),('E','liability'),('F','asset')]:annual_cache[f'{col}{index}']=a[key]
    notes=w.create_sheet('口径校验');notes.append(['项目','说明']);notes.append(['来源','西安公司使用权资产.xlsx；普源精电租赁合同.pdf'])
    for i,note in enumerate(result['notes'],1):notes.append([f'口径{i}',note])
    notes.append(['初始租赁负债',result['initial']['liability']]);notes.append(['初始使用权资产',result['initial']['asset']])
    notes.append(['期末资产余额',rows[-1]['closingAsset']]);notes.append(['期末负债余额',rows[-1]['closingLiability']])
    notes.append(['主表月租金合计与固定现金流差额',cents(sum(r['monthlyRent'] for r in rows)-result['initial']['totalFixed'])])
    phase_sheet=None;phase_cache={}
    if p.get('rentPhases'):
        phase_sheet=w.create_sheet('租金阶段');phase_sheet.append(['阶段','阶段开始日','阶段结束日','阶段月数','付款频率（月，0为整段一次）','付款时点','阶段含税租金总额','原表月租金','原表尾差','原约定支付期'])
        offset=0
        for index,phase in enumerate(p['rentPhases']):
            row=index+2;refs=[f"'租金参数'!D{i+5}" for i,st in enumerate(stages) if st.get('sourcePhase')==index]
            phase_sheet.append([index+1,shift_month(p['start'],offset),shift_month(p['start'],offset+phase['months'])-timedelta(days=1),phase['months'],phase['frequency'],{'advance':'期初付','arrears':'期末付','custom':'原约定期次'}[phase['timing']],
                                '=SUM('+','.join(refs)+')',phase.get('monthlyNet',0),phase.get('monthlyTail',0),phase.get('paymentPeriod') if phase['timing']=='custom' else None])
            phase_cache[f'G{row}']=phase['gross'];offset+=phase['months']
            phase_sheet[f'B{row}'].number_format=phase_sheet[f'C{row}'].number_format='yyyy/mm/dd'
    advanced_sheet=None
    if p.get('advanced'):
        advanced_sheet=w.create_sheet('高级参数');advanced_sheet.append(['高级条款参数','参数键','当前值'])
        for key,value in p['advanced'].items():
            advanced_sheet.append([{'freeEnabled':'存在免租期','freePeriods':'免租区间','substanceEnabled':'额外实质固定付款','variableType':'可变付款类型','renewalEnabled':'续租选择权','purchaseEnabled':'购买选择权','terminationEnabled':'提前终止选择权','directCosts':'初始直接费用','restorationPV':'复原义务现值','rateEvidence':'利率判断依据'}.get(key,key),key,json.dumps(value,ensure_ascii=False,allow_nan=False)])
    links=w.create_sheet('参数关联');links.append(['参数','参数位置','计算表位置','说明'])
    def link(sheet,address,target,description):
        sheet[address].hyperlink=Hyperlink(ref=address,location=target,tooltip=description)
        sheet[address].comment=Comment(description,'租赁计算系统')
    for address,param_cell in [('C1','B2'),('C2','B3'),('C3','B4'),('C4','B6'),('C6','B7'),('C8','B9'),('C10','B11'),('C11','B12')]:
        link(s,address,f"'计算参数'!{param_cell}",'查看／修改对应源参数；通过公式重算计算表。')
        link(ps,param_cell,f"'租赁计算明细'!{address}",'查看对应计算结果。')
        links.append([ps[f'A{ps[param_cell].row}'].value,'计算参数!'+param_cell,'租赁计算明细!'+address,'同一记录的参数与结果'])
    for index,st in enumerate(stages,5):
        payrow=18+st['paymentPeriod'];startrow=18+st['fromPeriod'];endrow=18+st['toPeriod']
        for col,dest in [('D',f'F{payrow}'),('I',f'F{payrow}'),('G',f'D{startrow}'),('H',f'D{endrow}')]:
            link(rs,f'{col}{index}',f"'租赁计算明细'!{dest}",'返回该阶段对应的计算结果。')
        links.append([f'第{index-4}阶段约定付款',f'租金参数!D{index} / I{index}',f'租赁计算明细!F{payrow}','付款额按支付期汇总；同一期多段付款时为合计'])
    for r in rows:
        row=r['period']+18;st_index=next(i for i,st in enumerate(stages) if st['fromPeriod']<=r['period']<=st['toPeriod'])+5
        payment_indices=[i+5 for i,st in enumerate(stages) if st['paymentPeriod']==r['period']]
        link(s,f'D{row}',f"'租金参数'!G{st_index}",'查看本期月租金及段末尾差来源。')
        link(s,f'J{row}',"'计算参数'!B9",'查看用于实际利率法摊销的年利率。')
        link(s,f'P{row}',"'计算参数'!B9",'利息为上期负债×年利率÷12；末笔利息按原表结清尾差。')
        target=f"'租金参数'!D{payment_indices[0]}" if payment_indices else "'租金参数'!I4"
        link(s,f'F{row}',target,'本期付款来源：'+('、'.join(f'租金参数第{i-4}阶段' for i in payment_indices) if payment_indices else '无对应支付期；查看各阶段约定支付期。'))
    first_zero=next((i+5 for i,st in enumerate(stages) if st['paymentPeriod']==0),None)
    if first_zero:link(s,'F18',f"'租金参数'!D{first_zero}",'查看确认时点付款对应的阶段。')
    for sheet in (ps,rs,ys,notes,links)+((advanced_sheet,) if advanced_sheet else ())+((phase_sheet,) if phase_sheet else ()):
        sheet.sheet_view.showGridLines=False;sheet.freeze_panes='B2'
        for row in sheet:
            for c in row:
                c.font=Font(name='微软雅黑',size=11,color='263340');c.alignment=Alignment(vertical='center',wrap_text=True)
                c.border=Border(bottom=side)
                if isinstance(c.value,(int,float)):c.number_format='#,##0.00;[Red](#,##0.00);–'
                if c.row==1:c.fill=blue;c.font=Font(name='微软雅黑',size=11,bold=True,color='263340')
        for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=24
    ps.column_dimensions['B'].width=100;notes.column_dimensions['B'].width=110
    for row in notes:notes.row_dimensions[row[0].row].height=45
    ps['B9'].number_format=ps['B10'].number_format='0.00%'
    ps['B8'].number_format=ps['B18'].number_format='0'
    for i in range(5,5+len(stages)):rs[f'E{i}'].number_format='0.00%'
    for i in range(2,2+len(result['annual'])):ys[f'A{i}'].number_format='0'
    w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    buffer=io.BytesIO();w.save(buffer)
    cache_sheets={'xl/worksheets/sheet1.xml':cache,'xl/worksheets/sheet3.xml':rent_cache,'xl/worksheets/sheet4.xml':annual_cache}
    if phase_sheet:cache_sheets[f'xl/worksheets/sheet{w.worksheets.index(phase_sheet)+1}.xml']=phase_cache
    return _cache_xlsx(buffer.getvalue(),cache_sheets)


if __name__ == '__main__':
    result=calculate(baseline())
    report=audit_baseline(result)
    print(json.dumps({'initial':result['initial'],'totals':result['totals'],'audit':report},ensure_ascii=False,indent=2))
    if not report['matched']:raise SystemExit(1)
    out=ROOT/'outputs';out.mkdir(exist_ok=True)
    (out/'西安租赁计算表-系统生成.xlsx').write_bytes(export_excel(result))
