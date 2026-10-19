"""Recover parameters from this system's export or the supplied DTT layout.

Read formulas and their saved values; never execute a workbook or infer missing
contract rights. The main schedule is recalculated to verify the recovered data.
"""
import hashlib
import io
import json
import math
import re
import zipfile
from collections import Counter
from datetime import date, datetime

import openpyxl

from lease_engine import calculate, cents, COLS


def iso(value):
    if isinstance(value,(datetime,date)):return value.isoformat()[:10]
    text=str(value or '')[:10]
    if text:date.fromisoformat(text)
    return text


def number(value, label):
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
        raise ValueError(f'{label}缺少可读取的数值，请用Excel计算并保存文件后重试。')
    return value


def comparison(result, formula_sheet, value_sheet):
    checked=0;differences=[];missing=0
    expected_rows={int(value_sheet.cell(r,1).value):r for r in range(19,value_sheet.max_row+1)
                   if isinstance(value_sheet.cell(r,1).value,(int,float)) and 1<=value_sheet.cell(r,1).value<=result['params']['months']}
    for row in result['rows']:
        original_row=expected_rows.get(row['period'])
        if original_row is None:missing+=len(COLS);continue
        for col,key in COLS.items():
            cell=f'{col}{original_row}';value=value_sheet[cell].value
            if value is None and formula_sheet[cell].data_type=='f':missing+=1;continue
            value=0 if value is None else value
            checked+=1
            if not isinstance(value,(int,float)) or abs(value-row[key])>(1e-12 if key in ('annualRate','pvFactor') else .00001):
                differences.append(dict(cell=cell,original=value,recalculated=row[key]))
    for cell,key in [('R18','liability'),('T18','asset'),('F18','upfront'),('C7','totalFixed')]:
        value=value_sheet[cell].value
        if value is None:missing+=1;continue
        checked+=1
        if not isinstance(value,(int,float)) or abs(value-result['initial'][key])>.00001:
            differences.append(dict(cell=cell,original=value,recalculated=result['initial'][key]))
    return dict(checked=checked,missing=missing,differenceCount=len(differences),differences=differences[:20],matched=checked>0 and not missing and not differences)


def reverse_excel(blob):
    if len(blob)>10_000_000:raise ValueError('计算表不能超过10MB。')
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            if sum(i.file_size for i in z.infolist())>50_000_000:raise ValueError('文件解压后过大，请使用单份租赁计算表。')
        formulas=openpyxl.load_workbook(io.BytesIO(blob),data_only=False,keep_links=False)
        values=openpyxl.load_workbook(io.BytesIO(blob),data_only=True,keep_links=False)
    except (zipfile.BadZipFile,KeyError):raise ValueError('请上传有效的.xlsx计算表。')
    digest=hashlib.sha256(blob).hexdigest()[:12]
    trace=[];warnings=[]
    if '计算参数' in values.sheetnames and '租金参数' in values.sheetnames and '租赁计算明细' in values.sheetnames:
        s=values['计算参数'];rs=values['租金参数'];rf=formulas['租金参数'];main=values['租赁计算明细'];mf=formulas['租赁计算明细']
        p=dict(company=str(s['B2'].value or ''),name=str(s['B3'].value or ''),lessor=str(s['B4'].value or ''),
               availableDate=iso(s['B5'].value),start=iso(s['B6'].value),end=iso(s['B7'].value),months=int(number(s['B8'].value,'计算期数')),
               annualRate=number(s['B9'].value,'年利率'),taxRate=number(s['B10'].value,'税率'),area=s['B11'].value or 0,
               deposit=s['B12'].value or 0,assetAdjustment=s['B13'].value or 0,freeEnd=iso(s['B14'].value),policy=str(s['B15'].value or ''),
               displayMode='source' if str(s['B16'].value).startswith('保留') else 'auto',
               recordId=str(s['B17'].value or 'IMPORT-'+digest),revision=int(s['B18'].value or 1),stages=[])
        for i in range(5,rs.max_row+1):
            if rs.cell(i,1).value is None:continue
            count=int(number(rs.cell(i,3).value,'阶段结束期')-number(rs.cell(i,2).value,'阶段开始期')+1)
            gross=number(rs.cell(i,4).value,'含税阶段租金');period=int(number(rs.cell(i,9).value,'约定支付期'))
            if rf.cell(i,5).data_type!='f' and rs.cell(i,5).value!=p['taxRate']:
                raise ValueError('租金参数中的分段税率与主参数不同；当前模型需先统一税率。')
            monthly_col,tail_col=(10,11) if rs['J4'].value=='原表月租金输入' else (7,8)
            monthly=number(rs.cell(i,monthly_col).value,'月租金展示值') if p['displayMode']=='source' else 0
            tail=number(rs.cell(i,tail_col).value,'段末尾差') if p['displayMode']=='source' else 0
            p['stages'].append(dict(months=count,gross=gross,paymentPeriod=period,monthlyNet=monthly,monthlyTail=tail))
            trace.append(dict(parameter=f'第{len(p["stages"])}阶段付款额',value=gross,source=f'租金参数!D{i}；约定支付期 I{i}'))
        for name,cell in [('租赁开始日','B6'),('租赁结束日','B7'),('计算期数','B8'),('年利率','B9'),('税率','B10')]:
            trace.append(dict(parameter=name,value=str(s[cell].value),source='计算参数!'+cell))
        p['contractId']=p['recordId']
        p['freeStart']=iso(s['B20'].value) if s['B20'].value else (p['availableDate'] if p['freeEnd'] else '')
        if '高级参数' in values.sheetnames:
            p['advanced']={str(row[1]):json.loads(row[2]) for row in values['高级参数'].iter_rows(min_row=2,values_only=True) if row[1]}
        if '租金阶段' in values.sheetnames:
            p['rentPhases']=[]
            for row in values['租金阶段'].iter_rows(min_row=2,values_only=True):
                if not row[0]:continue
                phase=dict(months=int(number(row[3],'阶段月数')),frequency=int(number(row[4],'付款频率')),timing={'期初付':'advance','期末付':'arrears','原约定期次':'custom'}.get(row[5],''),gross=number(row[6],'阶段租金总额'),monthlyNet=number(row[7],'原表月租金'),monthlyTail=number(row[8],'原表尾差'))
                if phase['timing']=='custom':phase['paymentPeriod']=int(number(row[9],'原约定支付期'))
                p['rentPhases'].append(phase)
        kind='系统关联计算表'
    else:
        candidates=[sheet for sheet in values if sheet['B14'].value=='期间' and sheet['A19'].value==1 and sheet['R16'].value in ('支付后余额','支付后负债余额')]
        if len(candidates)!=1:raise ValueError('尚不能识别此表结构。目前支持系统导出的关联计算表，以及西安这类德勤原表。')
        main=candidates[0];mf=formulas[main.title]
        period_rows=[r for r in range(19,main.max_row+1) if isinstance(main.cell(r,1).value,int) and 1<=main.cell(r,1).value<=600]
        if not period_rows or [main.cell(r,1).value for r in period_rows]!=list(range(1,len(period_rows)+1)):
            raise ValueError('原表期数不连续，需先核对计算期间。')
        n=len(period_rows);start=iso(main.cell(period_rows[0],2).value);end=iso(main.cell(period_rows[-1],3).value)
        annual=number(main['J19'].value,'年利率')
        if any(abs(number(main.cell(r,10).value,'逐期年利率')-annual)>1e-12 for r in period_rows):
            raise ValueError('原表含多种利率，当前固定利率模型不能自动还原。')
        payments=[(0,18)] + [(int(main.cell(r,1).value),r) for r in period_rows if main.cell(r,6).value]
        payments=[(period,r) for period,r in payments if main.cell(r,6).value]
        if not payments or payments[0][0]!=0:raise ValueError('原表无法唯一识别首笔付款与覆盖期间，需要补充参数后核对。')
        p=dict(company=str(main['C1'].value or ''),name=str(main['C2'].value or ''),lessor=str(main['C3'].value or ''),
               availableDate='',freeEnd='',start=start,end=end,months=n,annualRate=annual,taxRate=0,
               area=float(re.search(r'[\d.]+',str(main['C10'].value)).group()) if re.search(r'[\d.]+',str(main['C10'].value)) else 0,
               deposit=main['C11'].value or 0,assetAdjustment=cents(number(main['T18'].value,'初始资产')-number(main['R18'].value,'初始负债')-number(main['F18'].value,'初始付款')),
               displayMode='source',recordId='IMPORT-'+digest,revision=1,stages=[],
               policy='从既有德勤计算表反查：以原表起算日及现金流复现；取得资产控制权日、初始免租期及未列示的选择权需结合合同补充。')
        taxes=[]
        for index,(period,row_index) in enumerate(payments):
            match=re.fullmatch(r'=ROUND\(\s*([\d.]+)\s*/\s*([\d.]+)\s*,\s*2\s*\)',str(mf.cell(row_index,6).value),re.I)
            if not match:raise ValueError(f'{main.title}!F{row_index}无法可靠识别含税金额及税率，请补充付款参数。')
            gross=float(match[1]);tax=round(float(match[2])-1,8);taxes.append(tax)
            next_period=payments[index+1][0] if index+1<len(payments) else n
            count=next_period-period
            if count<=0:raise ValueError('付款期与覆盖期无法唯一识别。')
            display=[number(main.cell(19+period+i,4).value,'月租金') for i in range(count)]
            base=Counter(display[:-1] or display).most_common(1)[0][0]
            if any(abs(v-base)>.00001 for v in display[:-1]):raise ValueError('同一付款阶段包含多种月租金，请先拆分租金阶段。')
            p['stages'].append(dict(months=count,gross=gross,paymentPeriod=period,monthlyNet=base,monthlyTail=cents(display[-1]-base)))
            trace.append(dict(parameter=f'第{index+1}阶段付款额／税率',value=gross,source=f'{main.title}!F{row_index}；月租金 D{19+period}:D{18+next_period}'))
        if len(set(taxes))!=1:raise ValueError('原表包含不同税率，需人工确认分段税额。')
        p['taxRate']=taxes[0]
        warnings.append('原表未列示装修免租期和取得资产控制权日，已留空；不能仅凭计算起算日推断，也不能据此判断选择权不存在。')
        if iso(main['C6'].value)!=end:warnings.append(f'原表表头到期日为{iso(main["C6"].value)}，与末期结束日{end}冲突；本次按末期计算行提取，并保留冲突提示。')
        trace.extend([dict(parameter='租赁开始日',value=start,source=f'{main.title}!B19'),dict(parameter='租赁结束日',value=end,source=f'{main.title}!C{period_rows[-1]}'),dict(parameter='年利率',value=annual,source=f'{main.title}!J19；C8'),dict(parameter='计算期数',value=n,source=f'{main.title}!A{period_rows[-1]}')])
        kind='德勤原表反查'
    result=calculate(p)
    check=comparison(result,mf,main)
    if check['missing']:warnings.append(f'原文件有{check["missing"]}项公式没有保存计算值，已从参数重算；这些项目未完成原值比对。')
    if check['differenceCount']:warnings.append(f'反查后重算与原文件有{check["differenceCount"]}项差异，请先复核，不应视为完全还原。')
    return dict(params=p,kind=kind,trace=trace,warnings=warnings,comparison=check)
